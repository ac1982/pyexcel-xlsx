import os
from io import BytesIO

from openpyxl import Workbook
from pyexcel_xlsx import get_data

from .nose_tools import eq_


def test_hidden_row():
    data = get_data(
        os.path.join("tests", "fixtures", "hidden.xlsx"),
        skip_hidden_row_and_column=True,
        library="pyexcel-xlsx",
    )
    expected = [[1, 2], [7, 9]]
    eq_(data["Sheet1"], expected)


def test_complex_hidden_sheets():
    data = get_data(
        os.path.join("tests", "fixtures", "complex_hidden_sheets.xlsx"),
        skip_hidden_row_and_column=True,
        library="pyexcel-xlsx",
    )
    expected = [[1, 3, 5, 7, 9], [31, 33, 35, 37, 39], [61, 63, 65, 67]]
    eq_(data["Sheet1"], expected)


def test_grouped_hidden_columns(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append(list(range(1, 31)))
    sheet.append(list(range(31, 61)))
    sheet.column_dimensions["B"].width = 13
    sheet.column_dimensions.group("B", "D", hidden=True)
    sheet.column_dimensions["Z"].width = 13
    sheet.column_dimensions.group("Z", "AB", hidden=True)
    sheet.column_dimensions["F"].width = 13
    sheet.column_dimensions["F"].hidden = True
    path = tmp_path / "grouped-hidden.xlsx"
    book.save(path)
    book.close()

    result = get_data(
        str(path), skip_hidden_row_and_column=True, library="pyexcel-xlsx"
    )
    visible = [i for i in range(1, 31) if i not in (2, 3, 4, 6, 26, 27, 28)]
    eq_(result["Sheet"], [visible, [i + 30 for i in visible]])
    result = get_data(
        str(path), skip_hidden_row_and_column=False, library="pyexcel-xlsx"
    )
    eq_(result["Sheet"], [list(range(1, 31)), list(range(31, 61))])


def test_visible_column_overrides_hidden_group(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet.append([1, 2, 3, 4, 5])
    sheet.column_dimensions["B"].width = 13
    sheet.column_dimensions.group("B", "D", hidden=True)
    sheet.column_dimensions["C"].width = 13
    sheet.column_dimensions["C"].hidden = False
    path = tmp_path / "visible-override.xlsx"
    book.save(path)
    book.close()

    eq_(get_data(str(path), library="pyexcel-xlsx")["Sheet"], [[1, 3, 5]])


def test_hidden_group_with_merged_cells():
    book = Workbook()
    sheet = book.active
    sheet.append(["visible", "hidden", None, None, "last"])
    sheet.merge_cells("B1:D2")
    sheet.column_dimensions["B"].width = 13
    sheet.column_dimensions.group("B", "D", hidden=True)
    stream = BytesIO()
    book.save(stream)
    book.close()
    result = get_data(
        file_content=stream.getvalue(),
        afile=None,
        file_type="xlsx",
        detect_merged_cells=True,
        library="pyexcel-xlsx",
    )
    eq_(result["Sheet"], [["visible", "last"], []])
